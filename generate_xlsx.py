import csv
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BASE_URL = "https://raw.githubusercontent.com/CerveraDev/mor-wpb-workbook-tp/claude/sales-talking-points-90REn/slide-thumbnails/"

wb = Workbook()
ws = wb.active
ws.title = "Sales Talking Points"

header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="1F3864")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin = Side(style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

cell_align = Alignment(horizontal="left", vertical="top", wrap_text=True)

with open("MOR_WPB_Sales_Talking_Points.csv", newline="", encoding="utf-8") as f:
    reader = csv.reader(f)
    rows = list(reader)

headers = ["Thumbnail"] + rows[0]
for col_idx, header in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = border

ws.row_dimensions[1].height = 30

for row_idx, row in enumerate(rows[1:], start=2):
    slide_num = row[0].strip()
    img_url = f'{BASE_URL}slide-{int(slide_num):02d}.jpg'
    formula = f'=IMAGE("{img_url}")'

    img_cell = ws.cell(row=row_idx, column=1, value=formula)
    img_cell.alignment = Alignment(horizontal="center", vertical="center")
    img_cell.border = border

    for col_idx, value in enumerate(row, start=2):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = cell_align
        cell.border = border

    ws.row_dimensions[row_idx].height = 100

ws.column_dimensions["A"].width = 25
ws.column_dimensions["B"].width = 35
ws.column_dimensions["C"].width = 30
ws.column_dimensions["D"].width = 60
ws.column_dimensions["E"].width = 50
ws.column_dimensions["F"].width = 50

ws.freeze_panes = "A2"

wb.save("MOR_WPB_Sales_Talking_Points.xlsx")
print("Done: MOR_WPB_Sales_Talking_Points.xlsx")
